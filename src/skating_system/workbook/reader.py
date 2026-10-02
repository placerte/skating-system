from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from skating_system.workbook.findings import Finding, Severity
from skating_system.workbook.models import (
    Competition,
    EntryAssignment,
    EventMetadata,
    JudgeAssignment,
    MarkValue,
    ScoreRow,
    ScoreSheet,
    WorkbookData,
)
from skating_system.workbook.schema import (
    COMPETITIONS_CONTRACT,
    CORE_SHEETS,
    ENTRIES_CONTRACT,
    EVENT_CONTRACT,
    JUDGES_CONTRACT,
    SCHEMA_VERSION,
    SheetContract,
    comparison_key,
    normalize_header,
    normalize_text,
    score_sheet_names,
)


@dataclass(frozen=True)
class ReadResult:
    """Parsed workbook data and non-fatal/fatal parsing findings."""

    data: WorkbookData | None
    findings: tuple[Finding, ...]


def read_workbook(
    path: Path,
    *,
    require_score_sheets: bool = True,
    read_score_sheets: bool = True,
) -> ReadResult:
    """Read an event workbook without modifying it."""

    findings: list[Finding] = []
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (OSError, BadZipFile, InvalidFileException, ValueError) as exc:
        return ReadResult(
            data=None,
            findings=(Finding(Severity.ERROR, f"Could not open workbook: {exc}"),),
        )

    missing_sheets = [
        contract.name
        for contract in CORE_SHEETS
        if contract.name not in workbook.sheetnames
    ]
    for sheet_name in missing_sheets:
        findings.append(
            Finding(
                Severity.ERROR,
                f'Required sheet "{sheet_name}" is missing.',
                sheet=sheet_name,
            )
        )
    if missing_sheets:
        workbook.close()
        return ReadResult(data=None, findings=tuple(findings))

    event_columns = _read_columns(
        workbook[EVENT_CONTRACT.name], EVENT_CONTRACT, findings
    )
    competition_columns = _read_columns(
        workbook[COMPETITIONS_CONTRACT.name], COMPETITIONS_CONTRACT, findings
    )
    judge_columns = _read_columns(
        workbook[JUDGES_CONTRACT.name], JUDGES_CONTRACT, findings
    )
    entry_columns = _read_columns(
        workbook[ENTRIES_CONTRACT.name], ENTRIES_CONTRACT, findings
    )

    metadata = _read_metadata(workbook[EVENT_CONTRACT.name], event_columns, findings)
    competitions = _read_competitions(
        workbook[COMPETITIONS_CONTRACT.name], competition_columns, findings
    )
    judges = _read_judges(workbook[JUDGES_CONTRACT.name], judge_columns, findings)
    entries = _read_entries(workbook[ENTRIES_CONTRACT.name], entry_columns, findings)
    judges = _order_judges(judges, competitions)
    entries = _order_entries(entries, competitions)
    score_sheets: dict[str, ScoreSheet] = {}
    if read_score_sheets:
        score_sheets = _read_score_sheets(
            workbook,
            competitions,
            findings,
            require_score_sheets=require_score_sheets,
        )
    workbook.close()

    data = WorkbookData(
        path=path,
        metadata=metadata,
        competitions=tuple(competitions),
        judges=tuple(judges),
        entries=tuple(entries),
        score_sheets=score_sheets,
    )
    return ReadResult(data=data, findings=tuple(findings))


def _read_columns(
    sheet: Worksheet,
    contract: SheetContract,
    findings: list[Finding],
) -> dict[str, int]:
    columns: dict[str, int] = {}
    for cell in sheet[1]:
        header = normalize_header(cell.value)
        if not header:
            continue
        if header in columns:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'Duplicate column "{header}".',
                    sheet=sheet.title,
                    cell=cell.coordinate,
                )
            )
            continue
        if isinstance(cell.column, int):
            columns[header] = cell.column

    for required in contract.required_columns:
        if required not in columns:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'Required column "{required}" is missing.',
                    sheet=sheet.title,
                )
            )
    for header in columns:
        if header not in contract.columns:
            findings.append(
                Finding(
                    Severity.WARNING,
                    f'Unknown column "{header}" will be ignored.',
                    sheet=sheet.title,
                    cell=sheet.cell(1, columns[header]).coordinate,
                )
            )
    return columns


def _read_metadata(
    sheet: Worksheet,
    columns: dict[str, int],
    findings: list[Finding],
) -> EventMetadata:
    values: dict[str, str] = {}
    if "field" in columns and "value" in columns:
        for row_number in range(2, sheet.max_row + 1):
            field_cell = sheet.cell(row_number, columns["field"])
            field = normalize_header(field_cell.value)
            if not field:
                continue
            if field in values:
                findings.append(
                    Finding(
                        Severity.ERROR,
                        f'Duplicate event field "{field}".',
                        sheet=sheet.title,
                        cell=field_cell.coordinate,
                    )
                )
                continue
            values[field] = normalize_text(
                sheet.cell(row_number, columns["value"]).value
            )

    for field in ("schema_version", "event_name"):
        if not values.get(field):
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'Required event field "{field}" is missing or blank.',
                    sheet=sheet.title,
                )
            )
    version = values.get("schema_version", "")
    if version and version != SCHEMA_VERSION:
        findings.append(
            Finding(
                Severity.ERROR,
                f'Unsupported schema version "{version}"; expected "{SCHEMA_VERSION}".',
                sheet=sheet.title,
            )
        )
    return EventMetadata(
        schema_version=version, event_name=values.get("event_name", "")
    )


def _read_competitions(
    sheet: Worksheet,
    columns: dict[str, int],
    findings: list[Finding],
) -> list[Competition]:
    if not {"competition", "scoring_method"}.issubset(columns):
        return []
    competitions: list[Competition] = []
    for row_number in range(2, sheet.max_row + 1):
        if _row_is_blank(sheet, row_number):
            continue
        name_cell = sheet.cell(row_number, columns["competition"])
        method_cell = sheet.cell(row_number, columns["scoring_method"])
        name = normalize_text(name_cell.value)
        method = comparison_key(method_cell.value)
        alternate = _read_boolean(
            sheet,
            row_number,
            columns.get("alternate_enabled"),
            findings,
            competition=name,
        )
        competitions.append(
            Competition(
                name=name,
                scoring_method=method,
                alternate_enabled=alternate,
                status=_optional_text(sheet, row_number, columns.get("status")),
                notes=_optional_text(sheet, row_number, columns.get("notes")),
                callback_advance_count=_read_positive_integer(
                    sheet,
                    row_number,
                    columns.get("callback_advance_count"),
                    findings,
                    competition=name,
                    field="callback_advance_count",
                ),
                source_row=row_number,
            )
        )
        if not name:
            findings.append(
                Finding(
                    Severity.ERROR,
                    "Competition name is blank.",
                    sheet=sheet.title,
                    cell=name_cell.coordinate,
                )
            )
        if not method:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'{name or "Competition row"}: scoring method is blank.',
                    sheet=sheet.title,
                    cell=method_cell.coordinate,
                    competition=name or None,
                )
            )
    return competitions


def _read_judges(
    sheet: Worksheet,
    columns: dict[str, int],
    findings: list[Finding],
) -> list[JudgeAssignment]:
    if not {"competition", "judge"}.issubset(columns):
        return []
    assignments: list[JudgeAssignment] = []
    for row_number in range(2, sheet.max_row + 1):
        if _row_is_blank(sheet, row_number):
            continue
        competition = normalize_text(
            sheet.cell(row_number, columns["competition"]).value
        )
        judge_cell = sheet.cell(row_number, columns["judge"])
        judge = normalize_text(judge_cell.value)
        assignments.append(
            JudgeAssignment(
                competition=competition,
                judge=judge,
                order=_read_order(
                    sheet,
                    row_number,
                    columns.get("order"),
                    findings,
                    competition=competition,
                    noun="judge",
                ),
                source_row=row_number,
            )
        )
        if not judge:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'{competition or "Competition row"}: judge name is blank.',
                    sheet=sheet.title,
                    cell=judge_cell.coordinate,
                    competition=competition or None,
                )
            )
    return assignments


def _read_entries(
    sheet: Worksheet,
    columns: dict[str, int],
    findings: list[Finding],
) -> list[EntryAssignment]:
    if not {"competition", "entry"}.issubset(columns):
        return []
    assignments: list[EntryAssignment] = []
    for row_number in range(2, sheet.max_row + 1):
        if _row_is_blank(sheet, row_number):
            continue
        competition = normalize_text(
            sheet.cell(row_number, columns["competition"]).value
        )
        entry_cell = sheet.cell(row_number, columns["entry"])
        entry = normalize_text(entry_cell.value)
        assignments.append(
            EntryAssignment(
                competition=competition,
                entry=entry,
                number=_optional_text(sheet, row_number, columns.get("number")),
                order=_read_order(
                    sheet,
                    row_number,
                    columns.get("order"),
                    findings,
                    competition=competition,
                    noun="entry",
                ),
                source_row=row_number,
            )
        )
        if not entry:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'{competition or "Competition row"}: entry name is blank.',
                    sheet=sheet.title,
                    cell=entry_cell.coordinate,
                    competition=competition or None,
                )
            )
    return assignments


def _read_score_sheets(
    workbook: Workbook,
    competitions: list[Competition],
    findings: list[Finding],
    *,
    require_score_sheets: bool,
) -> dict[str, ScoreSheet]:
    names = score_sheet_names(
        [competition.name for competition in competitions if competition.name]
    )
    score_sheets: dict[str, ScoreSheet] = {}
    for competition in competitions:
        if not competition.name:
            continue
        sheet_name = names[competition.name]
        if sheet_name not in workbook.sheetnames:
            if require_score_sheets:
                findings.append(
                    Finding(
                        Severity.ERROR,
                        f'{competition.name}: score sheet "{sheet_name}" is missing.',
                        sheet=sheet_name,
                        competition=competition.name,
                    )
                )
            continue
        sheet = workbook[sheet_name]
        headers = [normalize_text(cell.value) for cell in sheet[1]]
        normalized = [normalize_header(header) for header in headers]
        if len(headers) < 2 or normalized[:2] != ["entry", "number"]:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f'{competition.name}: score sheet must start with "Entry" and "Number" columns.',
                    sheet=sheet_name,
                    competition=competition.name,
                )
            )
            continue
        judges = tuple(headers[2:])
        rows: list[ScoreRow] = []
        for row_number in range(2, sheet.max_row + 1):
            if _row_is_blank(sheet, row_number):
                continue
            marks = {
                judge: MarkValue(
                    value=sheet.cell(row_number, column_number).value,
                    cell=sheet.cell(row_number, column_number).coordinate,
                )
                for column_number, judge in enumerate(judges, start=3)
            }
            rows.append(
                ScoreRow(
                    entry=normalize_text(sheet.cell(row_number, 1).value),
                    number=normalize_text(sheet.cell(row_number, 2).value),
                    marks=marks,
                    source_row=row_number,
                )
            )
        score_sheets[comparison_key(competition.name)] = ScoreSheet(
            competition=competition.name,
            sheet_name=sheet_name,
            judges=judges,
            rows=tuple(rows),
        )
    return score_sheets


def _read_boolean(
    sheet: Worksheet,
    row_number: int,
    column_number: int | None,
    findings: list[Finding],
    *,
    competition: str,
) -> bool:
    if column_number is None:
        return False
    cell = sheet.cell(row_number, column_number)
    if cell.value is None or normalize_text(cell.value) == "":
        return False
    if isinstance(cell.value, bool):
        return cell.value
    value = comparison_key(cell.value)
    if value in {"true", "yes", "1", "=true()"}:
        return True
    if value in {"false", "no", "0", "=false()"}:
        return False
    findings.append(
        Finding(
            Severity.ERROR,
            f'{competition or "Competition row"}: alternate_enabled must be TRUE or FALSE.',
            sheet=sheet.title,
            cell=cell.coordinate,
            competition=competition or None,
        )
    )
    return False


def _read_order(
    sheet: Worksheet,
    row_number: int,
    column_number: int | None,
    findings: list[Finding],
    *,
    competition: str,
    noun: str,
) -> int | None:
    if column_number is None:
        return None
    value = sheet.cell(row_number, column_number).value
    if isinstance(value, bool) or value is None or normalize_text(value) == "":
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = normalize_text(value)
    if text.isdigit():
        return int(text)
    findings.append(
        Finding(
            Severity.WARNING,
            f"{competition}: {noun} order must be a positive integer; worksheet row order will be used.",
            sheet=sheet.title,
            cell=sheet.cell(row_number, column_number).coordinate,
            competition=competition,
        )
    )
    return None


def _read_positive_integer(
    sheet: Worksheet,
    row_number: int,
    column_number: int | None,
    findings: list[Finding],
    *,
    competition: str,
    field: str,
) -> int | None:
    if column_number is None:
        return None
    cell = sheet.cell(row_number, column_number)
    value = cell.value
    if value is None or normalize_text(value) == "":
        return None
    if isinstance(value, bool):
        parsed = None
    elif isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    else:
        text = normalize_text(value)
        parsed = int(text) if text.isdigit() else None
    if parsed is not None and parsed > 0:
        return parsed
    findings.append(
        Finding(
            Severity.ERROR,
            f"{competition or 'Competition row'}: {field} must be a positive integer.",
            sheet=sheet.title,
            cell=cell.coordinate,
            competition=competition or None,
        )
    )
    return None


def _optional_text(sheet: Worksheet, row_number: int, column_number: int | None) -> str:
    if column_number is None:
        return ""
    return normalize_text(sheet.cell(row_number, column_number).value)


def _row_is_blank(sheet: Worksheet, row_number: int) -> bool:
    return all(normalize_text(cell.value) == "" for cell in sheet[row_number])


def _order_judges(
    assignments: list[JudgeAssignment],
    competitions: list[Competition],
) -> list[JudgeAssignment]:
    result: list[JudgeAssignment] = []
    handled: set[str] = set()
    for competition in competitions:
        key = comparison_key(competition.name)
        group = [
            item for item in assignments if comparison_key(item.competition) == key
        ]
        orders = [item.order for item in group]
        valid = (
            bool(group)
            and all(order is not None and order > 0 for order in orders)
            and len(orders) == len(set(orders))
        )
        if valid:
            group.sort(key=lambda item: item.order if item.order is not None else 0)
        result.extend(group)
        handled.add(key)
    result.extend(
        item for item in assignments if comparison_key(item.competition) not in handled
    )
    return result


def _order_entries(
    assignments: list[EntryAssignment],
    competitions: list[Competition],
) -> list[EntryAssignment]:
    result: list[EntryAssignment] = []
    handled: set[str] = set()
    for competition in competitions:
        key = comparison_key(competition.name)
        group = [
            item for item in assignments if comparison_key(item.competition) == key
        ]
        orders = [item.order for item in group]
        valid = (
            bool(group)
            and all(order is not None and order > 0 for order in orders)
            and len(orders) == len(set(orders))
        )
        if valid:
            group.sort(key=lambda item: item.order if item.order is not None else 0)
        result.extend(group)
        handled.add(key)
    result.extend(
        item for item in assignments if comparison_key(item.competition) not in handled
    )
    return result
