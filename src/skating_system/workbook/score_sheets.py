from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from skating_system.workbook.findings import Finding, Severity
from skating_system.workbook.models import (
    Competition,
    EntryAssignment,
    JudgeAssignment,
    WorkbookData,
)
from skating_system.workbook.reader import read_workbook
from skating_system.workbook.schema import (
    SCORE_SHEET_PREFIX,
    SCORING_METHOD_CALLBACK,
    SCORING_METHOD_SKATING,
    comparison_key,
    normalize_header,
    normalize_text,
    score_sheet_names,
)
from skating_system.workbook.validation import validate_setup_data


@dataclass(frozen=True)
class BuildResult:
    """Outcome of creating or refreshing generated score sheets."""

    findings: tuple[Finding, ...]
    mappings: tuple[tuple[str, str], ...]
    changed: bool = False
    snapshot_path: Path | None = None

    @property
    def has_errors(self) -> bool:
        return any(finding.severity is Severity.ERROR for finding in self.findings)


@dataclass(frozen=True)
class _SheetPlan:
    competition: Competition
    sheet_name: str
    action: str
    preserved_marks: dict[tuple[tuple[str, str], str], object]
    destructive_reason: str | None = None


# GitHub issue #4: safely build and refresh generated score-entry sheets.
def build_score_sheets(path: Path, *, rebuild: bool = False) -> BuildResult:
    """Create or safely refresh generated score sheets in an event workbook."""

    read_result = read_workbook(
        path,
        require_score_sheets=False,
        read_score_sheets=False,
    )
    findings = list(read_result.findings)
    if read_result.data is None:
        return BuildResult(findings=tuple(findings), mappings=())

    data = read_result.data
    setup_findings, _competition_lookup = validate_setup_data(data)
    findings.extend(setup_findings)
    mappings = _sheet_mappings(data)
    if _has_errors(findings):
        return BuildResult(findings=tuple(findings), mappings=mappings)

    workbook = load_workbook(path, data_only=False)
    plans = [
        _plan_sheet(workbook, data, competition, sheet_name)
        for competition, sheet_name in _mapped_competitions(data)
    ]
    _warn_about_orphan_sheets(workbook, mappings, findings)
    destructive_plans = [plan for plan in plans if plan.destructive_reason]
    if destructive_plans and not rebuild:
        for plan in destructive_plans:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f"{plan.competition.name}: {plan.destructive_reason} "
                    "Run build-sheets with --rebuild to replace it after "
                    "creating a snapshot.",
                    sheet=plan.sheet_name,
                    competition=plan.competition.name,
                )
            )
        workbook.close()
        return BuildResult(findings=tuple(findings), mappings=mappings)

    snapshot_path: Path | None = None
    if destructive_plans:
        try:
            snapshot_path = _create_snapshot(path)
        except OSError as exc:
            findings.append(
                Finding(
                    Severity.ERROR,
                    f"Could not create required workbook snapshot: {exc}",
                )
            )
            workbook.close()
            return BuildResult(findings=tuple(findings), mappings=mappings)

    changed = False
    for plan in plans:
        if plan.action == "unchanged":
            continue
        if plan.sheet_name in workbook.sheetnames:
            del workbook[plan.sheet_name]
        sheet = workbook.create_sheet(plan.sheet_name)
        _write_score_sheet(sheet, data, plan.competition, plan.preserved_marks)
        changed = True

    if changed:
        try:
            _save_workbook(workbook, path)
        except OSError as exc:
            findings.append(Finding(Severity.ERROR, f"Could not save workbook: {exc}"))
        else:
            findings.append(
                Finding(Severity.INFO, "Score sheets were updated successfully.")
            )
    else:
        findings.append(Finding(Severity.INFO, "Score sheets are already current."))
    workbook.close()
    return BuildResult(
        findings=tuple(findings),
        mappings=mappings,
        changed=changed and not _has_errors(findings),
        snapshot_path=snapshot_path,
    )


def _sheet_mappings(data: WorkbookData) -> tuple[tuple[str, str], ...]:
    return tuple(
        (competition.name, sheet_name)
        for competition, sheet_name in _mapped_competitions(data)
    )


def _mapped_competitions(data: WorkbookData) -> list[tuple[Competition, str]]:
    competitions = [
        competition for competition in data.competitions if competition.name
    ]
    names = score_sheet_names([competition.name for competition in competitions])
    return [(competition, names[competition.name]) for competition in competitions]


def _plan_sheet(
    workbook: Workbook,
    data: WorkbookData,
    competition: Competition,
    sheet_name: str,
) -> _SheetPlan:
    if sheet_name not in workbook.sheetnames:
        return _SheetPlan(competition, sheet_name, "create", {})

    sheet = workbook[sheet_name]
    judges = _competition_judges(data, competition)
    entries = _competition_entries(data, competition)
    expected_headers = ["Entry", "Number", *[judge.judge for judge in judges]]
    actual_headers = [normalize_text(cell.value) for cell in sheet[1]]
    expected_rows = [(entry.entry, entry.number) for entry in entries]
    actual_rows = [
        (
            normalize_text(sheet.cell(row, 1).value),
            normalize_text(sheet.cell(row, 2).value),
        )
        for row in range(2, sheet.max_row + 1)
        if any(normalize_text(cell.value) for cell in sheet[row])
    ]
    if (
        _same_headers(actual_headers, expected_headers)
        and _same_rows(actual_rows, expected_rows)
        and _generated_layout_is_current(sheet, competition, len(entries), len(judges))
    ):
        return _SheetPlan(competition, sheet_name, "unchanged", {})

    preserved, reason = _collect_preserved_marks(sheet, judges, entries)
    return _SheetPlan(
        competition,
        sheet_name,
        "rebuild" if reason else "refresh",
        preserved,
        destructive_reason=reason,
    )


def _collect_preserved_marks(
    sheet: Worksheet,
    judges: list[JudgeAssignment],
    entries: list[EntryAssignment],
) -> tuple[dict[tuple[tuple[str, str], str], object], str | None]:
    if sheet.max_column < 2:
        return (
            {},
            "the existing score sheet has no recognizable Entry/Number structure.",
        )
    first_headers = [normalize_header(sheet.cell(1, column).value) for column in (1, 2)]
    if first_headers != ["entry", "number"]:
        return (
            {},
            "the existing score sheet has no recognizable Entry/Number structure.",
        )

    expected_judges = {comparison_key(judge.judge) for judge in judges}
    expected_entries = {_entry_key(entry.entry, entry.number) for entry in entries}
    actual_judges = [
        normalize_text(sheet.cell(1, column).value)
        for column in range(3, sheet.max_column + 1)
    ]
    judge_keys = [comparison_key(judge) for judge in actual_judges]
    nonblank_judges = [key for key in judge_keys if key]
    if len(nonblank_judges) != len(set(nonblank_judges)):
        return {}, "the existing score sheet has duplicate judge columns."

    populated_rows: list[tuple[int, tuple[str, str]]] = []
    for row_number in range(2, sheet.max_row + 1):
        entry = normalize_text(sheet.cell(row_number, 1).value)
        number = normalize_text(sheet.cell(row_number, 2).value)
        if not entry and not number and not _row_has_marks(sheet, row_number):
            continue
        populated_rows.append((row_number, _entry_key(entry, number)))
    row_keys = [key for _row, key in populated_rows]
    if len(row_keys) != len(set(row_keys)):
        return {}, "the existing score sheet has duplicate or ambiguous entry rows."

    preserved: dict[tuple[tuple[str, str], str], object] = {}
    reason: str | None = None
    for row_number, entry_key in populated_rows:
        for column_number, judge_key in enumerate(judge_keys, start=3):
            value = sheet.cell(row_number, column_number).value
            if value is None or normalize_text(value) == "":
                continue
            if entry_key not in expected_entries:
                reason = "an unregistered entry row contains marks that cannot be reassigned safely."
                continue
            if not judge_key or judge_key not in expected_judges:
                reason = "an unknown judge column contains marks that cannot be reassigned safely."
                continue
            preserved[(entry_key, judge_key)] = value
    return preserved, reason


def _write_score_sheet(
    sheet: Worksheet,
    data: WorkbookData,
    competition: Competition,
    preserved_marks: dict[tuple[tuple[str, str], str], object],
) -> None:
    judges = _competition_judges(data, competition)
    entries = _competition_entries(data, competition)
    sheet.append(["Entry", "Number", *[judge.judge for judge in judges]])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for entry in entries:
        entry_key = _entry_key(entry.entry, entry.number)
        marks = [
            preserved_marks.get((entry_key, comparison_key(judge.judge)))
            for judge in judges
        ]
        sheet.append([entry.entry, entry.number, *marks])
    sheet.freeze_panes = "C2"
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 12
    for column_number in range(3, 3 + len(judges)):
        column_letter = get_column_letter(column_number)
        sheet.column_dimensions[column_letter].width = 18
    _add_mark_validation(sheet, competition, len(entries), len(judges))


def _add_mark_validation(
    sheet: Worksheet,
    competition: Competition,
    entry_count: int,
    judge_count: int,
) -> None:
    if entry_count == 0 or judge_count == 0:
        return
    if competition.scoring_method == SCORING_METHOD_SKATING:
        validation = DataValidation(
            type="whole",
            operator="between",
            formula1="1",
            formula2=str(entry_count),
            allow_blank=True,
        )
        validation.error = f"Enter a whole-number rank from 1 to {entry_count}."
    elif competition.scoring_method == SCORING_METHOD_CALLBACK:
        values = "Y,N,A" if competition.alternate_enabled else "Y,N"
        validation = DataValidation(
            type="list",
            formula1=f'"{values}"',
            allow_blank=True,
        )
        validation.error = f"Enter one of: {values}."
    else:
        return
    validation.errorTitle = "Invalid mark"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    start = sheet.cell(2, 3).coordinate
    end = sheet.cell(entry_count + 1, judge_count + 2).coordinate
    validation.add(f"{start}:{end}")


def _competition_judges(
    data: WorkbookData,
    competition: Competition,
) -> list[JudgeAssignment]:
    key = comparison_key(competition.name)
    return [judge for judge in data.judges if comparison_key(judge.competition) == key]


def _competition_entries(
    data: WorkbookData,
    competition: Competition,
) -> list[EntryAssignment]:
    key = comparison_key(competition.name)
    return [entry for entry in data.entries if comparison_key(entry.competition) == key]


def _entry_key(entry: str, number: str) -> tuple[str, str]:
    number_key = comparison_key(number)
    if number_key:
        return "number", number_key
    return "entry", comparison_key(entry)


def _same_headers(actual: list[str], expected: list[str]) -> bool:
    return [comparison_key(value) for value in actual] == [
        comparison_key(value) for value in expected
    ]


def _same_rows(actual: list[tuple[str, str]], expected: list[tuple[str, str]]) -> bool:
    return [_entry_key(*row) for row in actual] == [
        _entry_key(*row) for row in expected
    ]


def _row_has_marks(sheet: Worksheet, row_number: int) -> bool:
    return any(
        normalize_text(sheet.cell(row_number, column).value)
        for column in range(3, sheet.max_column + 1)
    )


def _generated_layout_is_current(
    sheet: Worksheet,
    competition: Competition,
    entry_count: int,
    judge_count: int,
) -> bool:
    if sheet.freeze_panes != "C2":
        return False
    validations = list(sheet.data_validations.dataValidation)
    if entry_count == 0 or judge_count == 0:
        return validations == []
    if len(validations) != 1:
        return False
    validation = validations[0]
    expected_range = f"C2:{sheet.cell(entry_count + 1, judge_count + 2).coordinate}"
    if str(validation.sqref) != expected_range:
        return False
    if competition.scoring_method == SCORING_METHOD_SKATING:
        return (
            validation.type == "whole"
            and validation.formula1 == "1"
            and validation.formula2 == str(entry_count)
        )
    if competition.scoring_method == SCORING_METHOD_CALLBACK:
        values = "Y,N,A" if competition.alternate_enabled else "Y,N"
        return validation.type == "list" and validation.formula1 == f'"{values}"'
    return False


def _warn_about_orphan_sheets(
    workbook: Workbook,
    mappings: tuple[tuple[str, str], ...],
    findings: list[Finding],
) -> None:
    expected = {sheet_name.casefold() for _competition, sheet_name in mappings}
    for sheet_name in workbook.sheetnames:
        if (
            sheet_name.startswith(SCORE_SHEET_PREFIX)
            and sheet_name.casefold() not in expected
        ):
            findings.append(
                Finding(
                    Severity.WARNING,
                    f'Orphan generated sheet "{sheet_name}" was left unchanged.',
                    sheet=sheet_name,
                )
            )


def _create_snapshot(path: Path) -> Path:
    timestamp = datetime.now().astimezone().strftime("%Y-%m-%d-%H%M%S")
    base = path.with_name(f"{path.stem}.snapshot-{timestamp}{path.suffix}")
    snapshot = base
    suffix = 2
    while snapshot.exists():
        snapshot = base.with_name(f"{base.stem}-{suffix}{base.suffix}")
        suffix += 1
    shutil.copy2(path, snapshot)
    return snapshot


def _save_workbook(workbook: Workbook, path: Path) -> None:
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.stem}.",
        suffix=path.suffix,
        dir=path.parent,
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        temporary_path.chmod(path.stat().st_mode)
        os.replace(temporary_path, path)
    except OSError:
        temporary_path.unlink(missing_ok=True)
        raise


def _has_errors(findings: list[Finding]) -> bool:
    return any(finding.severity is Severity.ERROR for finding in findings)
