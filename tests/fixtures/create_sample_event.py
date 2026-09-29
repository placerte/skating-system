from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from skating_system.workbook.schema import (
    COMPETITIONS_CONTRACT,
    ENTRIES_CONTRACT,
    EVENT_CONTRACT,
    JUDGES_CONTRACT,
    SCHEMA_VERSION,
)

FIXTURE_PATH = Path(__file__).with_name("sample_event.xlsx")


def main() -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    event = workbook.create_sheet(EVENT_CONTRACT.name)
    _append_table(
        event,
        EVENT_CONTRACT.columns,
        [
            ("schema_version", SCHEMA_VERSION),
            ("event_name", "Retro Boreal 2027"),
        ],
    )

    competitions = workbook.create_sheet(COMPETITIONS_CONTRACT.name)
    _append_table(
        competitions,
        COMPETITIONS_CONTRACT.columns,
        [
            ("Open Mix & Match", "skating", False, "setup", "Final"),
            ("Solo Jazz Prelims", "callback", True, "setup", "Preliminary"),
        ],
    )

    judges = workbook.create_sheet(JUDGES_CONTRACT.name)
    _append_table(
        judges,
        JUDGES_CONTRACT.columns,
        [
            ("Open Mix & Match", "Jane Smith", 1),
            ("Open Mix & Match", "Morgan Lee", 2),
            ("Open Mix & Match", "Sam Roy", 3),
            ("Solo Jazz Prelims", "Jane Smith", 1),
            ("Solo Jazz Prelims", "Morgan Lee", 2),
            ("Solo Jazz Prelims", "Sam Roy", 3),
        ],
    )

    entries = workbook.create_sheet(ENTRIES_CONTRACT.name)
    _append_table(
        entries,
        ENTRIES_CONTRACT.columns,
        [
            ("Open Mix & Match", "Alice Tremblay & Bob Smith", "101", 1),
            ("Open Mix & Match", "Carol Nguyen & Dan Roy", "102", 2),
            ("Solo Jazz Prelims", "Evelyn Chen", "201", 1),
            ("Solo Jazz Prelims", "Frank Martin", "202", 2),
        ],
    )

    workbook.save(FIXTURE_PATH)


def _append_table(
    sheet, columns: tuple[str, ...], rows: list[tuple[object, ...]]
) -> None:
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 40)


if __name__ == "__main__":
    main()
