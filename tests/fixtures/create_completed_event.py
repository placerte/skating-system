from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from skating_system.workbook.score_sheets import build_score_sheets

FIXTURE_DIRECTORY = Path(__file__).parent
SOURCE_PATH = FIXTURE_DIRECTORY / "sample_event.xlsx"
OUTPUT_PATH = FIXTURE_DIRECTORY / "completed_event.xlsx"


def main() -> None:
    copyfile(SOURCE_PATH, OUTPUT_PATH)
    result = build_score_sheets(OUTPUT_PATH)
    if result.has_errors:
        messages = "; ".join(finding.message for finding in result.findings)
        raise RuntimeError(messages)

    workbook = load_workbook(OUTPUT_PATH)
    skating = workbook["Score - Open Mix & Match"]
    skating_marks = ((1, 1, 1), (2, 2, 2))
    _write_marks(skating, skating_marks)

    callback = workbook["Score - Solo Jazz Prelims"]
    callback_marks = (("Y", "Y", "A"), ("N", "N", "N"))
    _write_marks(callback, callback_marks)

    workbook.save(OUTPUT_PATH)
    workbook.close()


def _write_marks(
    sheet: Worksheet,
    marks_by_entry: tuple[tuple[int | str, ...], ...],
) -> None:
    for row_offset, marks in enumerate(marks_by_entry, start=2):
        for column_offset, mark in enumerate(marks, start=3):
            sheet.cell(row_offset, column_offset, mark)


if __name__ == "__main__":
    main()
