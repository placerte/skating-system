from __future__ import annotations

import shutil
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from skating_system.cli import main
from skating_system.workbook.score_sheets import build_score_sheets

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "sample_event.xlsx"


def test_build_creates_both_layouts_and_reports_mappings(tmp_path: Path) -> None:
    path = _copy_fixture(tmp_path)

    result = build_score_sheets(path)

    assert not result.has_errors
    assert result.changed
    assert result.snapshot_path is None
    assert result.mappings == (
        ("Open Mix & Match", "Score - Open Mix & Match"),
        ("Solo Jazz Prelims", "Score - Solo Jazz Prelims"),
    )

    workbook = load_workbook(path)
    skating = workbook["Score - Open Mix & Match"]
    callback = workbook["Score - Solo Jazz Prelims"]
    assert [cell.value for cell in skating[1]] == [
        "Entry",
        "Number",
        "Jane Smith",
        "Morgan Lee",
        "Sam Roy",
    ]
    assert skating["A2"].value == "Alice Tremblay & Bob Smith"
    assert list(skating.data_validations.dataValidation)[0].type == "whole"
    callback_validation = list(callback.data_validations.dataValidation)[0]
    assert callback_validation.type == "list"
    assert callback_validation.formula1 == '"Y,N,A"'


def test_current_sheets_are_not_rewritten_or_snapshotted(tmp_path: Path) -> None:
    path = _copy_fixture(tmp_path)
    first = build_score_sheets(path)
    assert not first.has_errors
    before = sha256(path.read_bytes()).digest()

    second = build_score_sheets(path)

    assert not second.has_errors
    assert second.changed is False
    assert second.snapshot_path is None
    assert sha256(path.read_bytes()).digest() == before
    assert list(tmp_path.glob("*.snapshot-*.xlsx")) == []


def test_safe_refresh_reorders_and_preserves_marks_without_snapshot(
    tmp_path: Path,
) -> None:
    path = _copy_fixture(tmp_path)
    build_score_sheets(path)
    workbook = load_workbook(path)
    score = workbook["Score - Open Mix & Match"]
    score["C2"] = 1
    score["D2"] = 2
    judges = workbook["CompetitionJudges"]
    judges["C2"] = 2
    judges["C3"] = 1
    entries = workbook["CompetitionEntries"]
    entries["D2"] = 2
    entries["D3"] = 1
    workbook.save(path)

    result = build_score_sheets(path)

    assert not result.has_errors
    assert result.changed
    assert result.snapshot_path is None
    refreshed = load_workbook(path)["Score - Open Mix & Match"]
    assert [refreshed.cell(1, column).value for column in range(3, 6)] == [
        "Morgan Lee",
        "Jane Smith",
        "Sam Roy",
    ]
    assert refreshed["A2"].value == "Carol Nguyen & Dan Roy"
    assert refreshed["A3"].value == "Alice Tremblay & Bob Smith"
    assert refreshed["C3"].value == 2
    assert refreshed["D3"].value == 1


def test_unsafe_refresh_stops_without_writing_or_snapshot(tmp_path: Path) -> None:
    path = _workbook_with_unknown_mark(tmp_path)
    before = sha256(path.read_bytes()).digest()

    result = build_score_sheets(path)

    assert result.has_errors
    assert result.changed is False
    assert result.snapshot_path is None
    assert sha256(path.read_bytes()).digest() == before
    assert list(tmp_path.glob("*.snapshot-*.xlsx")) == []
    assert any("unknown judge column" in finding.message for finding in result.findings)


def test_rebuild_fails_safely_when_snapshot_cannot_be_created(
    tmp_path: Path,
    monkeypatch,
) -> None:
    path = _workbook_with_unknown_mark(tmp_path)
    before = sha256(path.read_bytes()).digest()

    def fail_copy(_source: Path, _target: Path) -> None:
        raise OSError("snapshot denied")

    monkeypatch.setattr(
        "skating_system.workbook.score_sheets.shutil.copy2",
        fail_copy,
    )

    result = build_score_sheets(path, rebuild=True)

    assert result.has_errors
    assert result.snapshot_path is None
    assert sha256(path.read_bytes()).digest() == before
    assert any("snapshot denied" in finding.message for finding in result.findings)


def test_rebuild_snapshots_then_drops_only_unmappable_marks(tmp_path: Path) -> None:
    path = _workbook_with_unknown_mark(tmp_path)
    before = path.read_bytes()

    result = build_score_sheets(path, rebuild=True)

    assert not result.has_errors
    assert result.changed
    assert result.snapshot_path is not None
    assert result.snapshot_path.read_bytes() == before
    refreshed = load_workbook(path)["Score - Open Mix & Match"]
    assert [cell.value for cell in refreshed[1]] == [
        "Entry",
        "Number",
        "Jane Smith",
        "Morgan Lee",
        "Sam Roy",
    ]
    assert refreshed["C2"].value == 1


def test_cli_prints_collision_mapping(tmp_path: Path, capsys) -> None:
    path = _copy_fixture(tmp_path)
    workbook = load_workbook(path)
    competitions = workbook["Competitions"]
    competitions["A2"] = "A very long competition name with extras"
    competitions["A3"] = "A very long competition name with extras!"
    for sheet_name in ("CompetitionJudges", "CompetitionEntries"):
        sheet = workbook[sheet_name]
        for row in range(2, sheet.max_row + 1):
            old_name = sheet.cell(row, 1).value
            sheet.cell(
                row,
                1,
                (
                    "A very long competition name with extras"
                    if old_name == "Open Mix & Match"
                    else "A very long competition name with extras!"
                ),
            )
    workbook.save(path)

    assert main(["build-sheets", str(path)]) == 0
    output = capsys.readouterr().out
    assert (
        '"A very long competition name with extras" -> "Score - A very long competition"'
        in output
    )
    assert (
        '"A very long competition name with extras!" -> "Score - A very long competit ~2"'
        in output
    )


def _copy_fixture(tmp_path: Path) -> Path:
    path = tmp_path / "event.xlsx"
    shutil.copy2(FIXTURE_PATH, path)
    return path


def _workbook_with_unknown_mark(tmp_path: Path) -> Path:
    path = _copy_fixture(tmp_path)
    build_score_sheets(path)
    workbook = load_workbook(path)
    score = workbook["Score - Open Mix & Match"]
    score["C2"] = 1
    score["F1"] = "Former Judge"
    score["F2"] = 2
    workbook.save(path)
    return path
