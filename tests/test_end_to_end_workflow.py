from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook

from skating_system.cli import main
from tests.pdf_helpers import inspect_pdf


FIXTURE = Path(__file__).parent / "fixtures" / "completed_event.xlsx"


def test_complete_competition_day_workflow(
    tmp_path: Path,
    capsys,
) -> None:
    workbook_path = tmp_path / "competition-day.xlsx"
    copyfile(FIXTURE, workbook_path)

    assert main(["build-sheets", str(workbook_path)]) == 0
    assert main(["validate", str(workbook_path)]) == 0
    assert main(["generate", "call-sheets", str(workbook_path)]) == 0
    assert main(["generate", "judge-cards", str(workbook_path)]) == 0
    assert main(["compute", str(workbook_path)]) == 0
    assert main(["report", "all", str(workbook_path)]) == 0

    output = capsys.readouterr().out
    assert 'MAP: "Open Mix & Match"' in output
    assert "Workbook structure and marks are valid" in output
    assert "PLACE: 1 | Alice Tremblay & Bob Smith" in output
    assert "CALLBACK: ADVANCE | 2.5 | Y=2 A=1 | Evelyn Chen" in output

    expected_reports = {
        "call-sheets.pdf",
        "judge-cards.pdf",
        "public-results.pdf",
        "management-results.pdf",
        "mc-results.pdf",
    }
    report_directory = tmp_path / "reports"
    assert {path.name for path in report_directory.glob("*.pdf")} == expected_reports
    for filename in expected_reports:
        inspection = inspect_pdf(report_directory / filename)
        assert inspection.page_count >= 1

    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook["Score - Open Mix & Match"].cell(2, 3).value == 1
    assert workbook["Score - Solo Jazz Prelims"].cell(2, 5).value == "A"
    workbook.close()
