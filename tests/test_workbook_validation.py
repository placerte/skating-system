from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from skating_system.cli import main
from skating_system.workbook.findings import Severity
from skating_system.workbook.reader import read_workbook
from skating_system.workbook.validation import (
    normalize_callback_mark,
    validate_workbook,
)


def test_valid_skating_workbook_is_ready_and_read_only(tmp_path: Path) -> None:
    path = _write_workbook(tmp_path)
    before = sha256(path.read_bytes()).digest()

    report = validate_workbook(path)

    assert report.has_errors is False
    assert report.findings[-1].severity is Severity.INFO
    assert sha256(path.read_bytes()).digest() == before


def test_validate_cli_returns_nonzero_for_missing_workbook(
    tmp_path: Path,
    capsys,
) -> None:
    missing = tmp_path / "missing.xlsx"

    assert main(["validate", str(missing)]) == 1
    assert "ERROR: Could not open workbook" in capsys.readouterr().out


def test_invalid_xlsx_file_is_reported_without_crashing(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    path.write_text("not an Excel workbook", encoding="utf-8")

    report = validate_workbook(path)

    assert report.has_errors
    assert report.findings[0].message.startswith("Could not open workbook")


def test_missing_core_sheet_is_an_error(tmp_path: Path) -> None:
    path = _write_workbook(tmp_path)
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    del workbook["CompetitionJudges"]
    workbook.save(path)

    report = validate_workbook(path)

    assert report.has_errors
    assert any("Required sheet" in finding.message for finding in report.findings)


def test_duplicates_and_broken_references_have_context(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path,
        competition_rows=[
            ("Open", "skating", False, "setup", ""),
            (" open ", "skating", False, "setup", ""),
        ],
        judge_rows=[("Missing", "Judge A", 1)],
        entry_rows=[
            ("Open", "Pair A", "101", 1),
            ("Open", " pair a ", "101", 2),
        ],
        include_score_sheet=False,
    )

    report = validate_workbook(path)
    messages = [finding.format() for finding in report.findings]

    assert any(
        "Duplicate competition" in message and "Competitions!A3" in message
        for message in messages
    )
    assert any(
        "referenced competition does not exist" in message for message in messages
    )
    assert any(
        message.startswith("WARNING") and "entry name" in message
        for message in messages
    )
    assert any(
        message.startswith("ERROR") and "entry number" in message
        for message in messages
    )


def test_complete_unique_order_controls_parsed_order(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path,
        judge_rows=[("Open", "Judge B", 2), ("Open", "Judge A", 1)],
        entry_rows=[
            ("Open", "Pair B", "102", 2),
            ("Open", "Pair A", "101", 1),
        ],
        score_headers=("Entry", "Number", "Judge A", "Judge B"),
        score_rows=[("Pair A", "101", 1, 1), ("Pair B", "102", 2, 2)],
    )

    result = read_workbook(path)

    assert result.data is not None
    assert [judge.judge for judge in result.data.judges] == ["Judge A", "Judge B"]
    assert [entry.entry for entry in result.data.entries] == ["Pair A", "Pair B"]


def test_incomplete_skating_marks_and_duplicate_ranks_are_errors(
    tmp_path: Path,
) -> None:
    path = _write_workbook(
        tmp_path,
        score_rows=[("Pair A", "101", 1), ("Pair B", "102", 1)],
    )

    report = validate_workbook(path)
    messages = [finding.message for finding in report.findings]

    assert report.has_errors
    assert any("ranked both entry" in message for message in messages)
    assert any("must use every rank" in message for message in messages)


def test_callback_aliases_are_normalized_and_disabled_alternate_is_rejected(
    tmp_path: Path,
) -> None:
    assert normalize_callback_mark(" yes ") == "Y"
    assert normalize_callback_mark("NO") == "N"
    assert normalize_callback_mark("Alt") == "A"

    path = _write_workbook(
        tmp_path,
        competition_rows=[("Callbacks", "callback", False, "setup", "")],
        judge_rows=[("Callbacks", "Judge A", 1)],
        entry_rows=[
            ("Callbacks", "Solo A", "201", 1),
            ("Callbacks", "Solo B", "202", 2),
        ],
        score_headers=("Entry", "Number", "Judge A"),
        score_rows=[("Solo A", "201", "yes"), ("Solo B", "202", "alt")],
        score_sheet_name="Score - Callbacks",
    )

    report = validate_workbook(path)

    assert any(
        "Alternate is disabled" in finding.message for finding in report.findings
    )


def test_excel_boolean_formulas_are_supported(tmp_path: Path) -> None:
    false_path = _write_workbook(
        tmp_path,
        competition_rows=[("Open", "skating", "=FALSE()", "setup", "")],
    )

    false_result = read_workbook(false_path)

    assert false_result.data is not None
    assert not false_result.findings
    assert false_result.data.competitions[0].alternate_enabled is False

    true_path = _write_workbook(
        tmp_path,
        competition_rows=[("Open", "skating", "=TRUE()", "setup", "")],
    )

    true_result = read_workbook(true_path)

    assert true_result.data is not None
    assert not true_result.findings
    assert true_result.data.competitions[0].alternate_enabled is True


def _write_workbook(
    tmp_path: Path,
    *,
    competition_rows: list[tuple[object, ...]] | None = None,
    judge_rows: list[tuple[object, ...]] | None = None,
    entry_rows: list[tuple[object, ...]] | None = None,
    score_headers: tuple[object, ...] = ("Entry", "Number", "Judge A"),
    score_rows: list[tuple[object, ...]] | None = None,
    score_sheet_name: str = "Score - Open",
    include_score_sheet: bool = True,
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_sheet(
        workbook,
        "Event",
        ("field", "value"),
        [("schema_version", "1"), ("event_name", "Test Event")],
    )
    _add_sheet(
        workbook,
        "Competitions",
        ("competition", "scoring_method", "alternate_enabled", "status", "notes"),
        competition_rows or [("Open", "skating", False, "setup", "")],
    )
    _add_sheet(
        workbook,
        "CompetitionJudges",
        ("competition", "judge", "order"),
        judge_rows or [("Open", "Judge A", 1)],
    )
    _add_sheet(
        workbook,
        "CompetitionEntries",
        ("competition", "entry", "number", "order"),
        entry_rows or [("Open", "Pair A", "101", 1), ("Open", "Pair B", "102", 2)],
    )
    if include_score_sheet:
        _add_sheet(
            workbook,
            score_sheet_name,
            score_headers,
            score_rows or [("Pair A", "101", 1), ("Pair B", "102", 2)],
        )
    path = tmp_path / "event.xlsx"
    workbook.save(path)
    return path


def _add_sheet(
    workbook: Workbook,
    name: str,
    headers: tuple[object, ...],
    rows: list[tuple[object, ...]],
) -> Worksheet:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    return sheet
