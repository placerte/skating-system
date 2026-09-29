from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from skating_system.workbook.schema import (
    COMPETITIONS_CONTRACT,
    CORE_SHEETS,
    ENTRIES_CONTRACT,
    EVENT_CONTRACT,
    JUDGES_CONTRACT,
    comparison_key,
    normalize_header,
    normalize_text,
    score_sheet_names,
)

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "sample_event.xlsx"


def test_core_contract_columns() -> None:
    assert [contract.name for contract in CORE_SHEETS] == [
        "Event",
        "Competitions",
        "CompetitionJudges",
        "CompetitionEntries",
    ]
    assert EVENT_CONTRACT.required_columns == ("field", "value")
    assert COMPETITIONS_CONTRACT.required_columns == (
        "competition",
        "scoring_method",
    )
    assert JUDGES_CONTRACT.required_columns == ("competition", "judge")
    assert ENTRIES_CONTRACT.required_columns == ("competition", "entry")


def test_text_normalization_preserves_display_case() -> None:
    assert normalize_text("  Open   Mix & Match ") == "Open Mix & Match"
    assert comparison_key(" OPEN mix & match ") == "open mix & match"
    assert normalize_header(" Scoring Method ") == "scoring_method"


def test_score_sheet_names_are_safe_and_deterministic() -> None:
    competition_names = [
        "A very long competition name with extras",
        "A very long competition name with extras!",
        "Finals: Open/Mix?Match",
    ]

    names = score_sheet_names(competition_names)

    assert names == {
        "A very long competition name with extras": "Score - A very long competition",
        "A very long competition name with extras!": "Score - A very long competit ~2",
        "Finals: Open/Mix?Match": "Score - Finals- Open-Mix-Match",
    }
    assert all(len(name) <= 31 for name in names.values())


def test_sample_workbook_matches_contract() -> None:
    workbook = load_workbook(FIXTURE_PATH, data_only=False)

    assert workbook.sheetnames == [contract.name for contract in CORE_SHEETS]
    assert _headers(workbook["Event"]) == EVENT_CONTRACT.columns
    assert _headers(workbook["Competitions"]) == COMPETITIONS_CONTRACT.columns
    assert _headers(workbook["CompetitionJudges"]) == JUDGES_CONTRACT.columns
    assert _headers(workbook["CompetitionEntries"]) == ENTRIES_CONTRACT.columns

    metadata = {
        row[0].value: str(row[1].value)
        for row in workbook["Event"].iter_rows(min_row=2)
    }
    assert metadata == {
        "schema_version": "1",
        "event_name": "Retro Boreal 2027",
    }

    methods = {
        row[1].value
        for row in workbook["Competitions"].iter_rows(min_row=2)
        if row[0].value is not None
    }
    assert methods == {"skating", "callback"}


def _headers(sheet) -> tuple[str, ...]:
    return tuple(cell.value for cell in sheet[1])
