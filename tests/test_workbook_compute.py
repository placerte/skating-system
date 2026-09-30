from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from shutil import copyfile

from openpyxl import load_workbook

from skating_system.cli import main
from skating_system.scoring.callback import CallbackResult
from skating_system.services.skating_scorer import SolveResult
from skating_system.services.workbook_compute import compute_workbook
from skating_system.workbook.score_sheets import build_score_sheets


FIXTURE = Path(__file__).parent / "fixtures" / "sample_event.xlsx"


def _completed_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "event.xlsx"
    copyfile(FIXTURE, path)
    result = build_score_sheets(path)
    assert not result.has_errors
    workbook = load_workbook(path)
    skating = workbook["Score - Open Mix & Match"]
    skating.cell(2, 3, 1)
    skating.cell(2, 4, 1)
    skating.cell(2, 5, 1)
    skating.cell(3, 3, 2)
    skating.cell(3, 4, 2)
    skating.cell(3, 5, 2)
    callback = workbook["Score - Solo Jazz Prelims"]
    callback.cell(2, 3, "Y")
    callback.cell(2, 4, "yes")
    callback.cell(2, 5, "A")
    callback.cell(3, 3, "N")
    callback.cell(3, 4, "no")
    callback.cell(3, 5, "N")
    workbook.save(path)
    workbook.close()
    return path


def test_compute_all_dispatches_to_both_engines_without_mutation(
    tmp_path: Path,
) -> None:
    path = _completed_workbook(tmp_path)
    before = sha256(path.read_bytes()).digest()

    result = compute_workbook(path)

    assert not result.has_errors
    assert len(result.competitions) == 2
    assert isinstance(result.competitions[0].result, SolveResult)
    assert isinstance(result.competitions[1].result, CallbackResult)
    callback = result.competitions[1].result
    assert len(callback.selection.advancing_entry_ids) == 1
    assert sha256(path.read_bytes()).digest() == before


def test_named_competition_filter_is_case_insensitive(tmp_path: Path) -> None:
    path = _completed_workbook(tmp_path)
    workbook = load_workbook(path)
    workbook["Score - Open Mix & Match"].cell(2, 3).value = None
    workbook.save(path)
    workbook.close()

    result = compute_workbook(path, " solo jazz PRELIMS ")

    assert not result.has_errors
    assert [item.competition.name for item in result.competitions] == [
        "Solo Jazz Prelims"
    ]


def test_unknown_competition_error_lists_available_names(tmp_path: Path) -> None:
    path = _completed_workbook(tmp_path)

    result = compute_workbook(path, "Mystery")

    assert result.has_errors
    message = " ".join(item.message for item in result.findings)
    assert "Mystery" in message
    assert "Open Mix & Match" in message


def test_invalid_workbook_cannot_produce_results(tmp_path: Path) -> None:
    path = _completed_workbook(tmp_path)
    workbook = load_workbook(path)
    workbook["Score - Open Mix & Match"].cell(2, 3).value = None
    workbook.save(path)
    workbook.close()

    result = compute_workbook(path)

    assert result.has_errors
    assert not result.competitions


def test_compute_cli_prints_callback_policy(
    tmp_path: Path,
    capsys,
) -> None:
    path = _completed_workbook(tmp_path)

    assert main(["compute", str(path), "--competition", "Solo Jazz Prelims"]) == 0

    output = capsys.readouterr().out
    assert "ADVANCE" in output
    assert "POLICY: Y=1 A=0.5 N=0" in output
